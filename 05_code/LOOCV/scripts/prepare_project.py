#!/usr/bin/env python3
"""Prepare immutable full-data gates and 48 nested_LOOCV outer-LOOCV cases."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import random
import shutil
from pathlib import Path

from openpyxl import load_workbook

from workflow import (
    OPS,
    SCREENING_FIELDS,
    mapping_rows,
    ordered_features,
    read_dicts,
    screen,
    write_dicts,
    write_sisso_inputs,
)


ROOT = Path(__file__).resolve().parents[1]
FIXED_SEED = 20260902
SOURCES = {
    "mono": {
        "file": ROOT / "input" / "mono_LP_56_2.xlsx",
        "sha256": "89D6A3C7EF88AC0934D6ACAB34738802FEF96058FA8A72BE706CE03BA354EF0B",
        "id_column": 2,
        "source_name_column": 1,
        "feature_start": 3,
        "feature_end": 58,
        "target_column": 59,
        "prefix": "LP",
        "target": "raw",
        "expected_selected": 28,
        "reference": ROOT / "reference" / "mono",
    },
    "bi": {
        "file": ROOT / "input" / "bi_LPP_59.xlsx",
        "sha256": "94F2CE19FC5B1B7DE32A07D5085274BFA10563844596951DDC3B3F82FD3EC5A0",
        "id_column": 1,
        "source_name_column": 1,
        "feature_start": 2,
        "feature_end": 60,
        "target_column": 61,
        "prefix": "LPP",
        "target": "log",
        "expected_selected": 31,
        "reference": ROOT / "reference" / "bi",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def reset_directory(path: Path) -> None:
    resolved = path.resolve()
    if resolved.parent != ROOT or resolved.name not in {"cases", "full_data", "truth", "results"}:
        raise RuntimeError(f"Refusing to reset unexpected directory: {resolved}")
    if resolved.exists():
        shutil.rmtree(resolved)
    resolved.mkdir(parents=True)


def load_source(kind: str) -> tuple[list[dict[str, object]], list[str], str]:
    spec = SOURCES[kind]
    path = Path(spec["file"])
    actual_hash = sha256(path)
    if actual_hash != spec["sha256"]:
        raise RuntimeError(f"{path.name}: SHA256 mismatch: {actual_hash}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    columns = range(int(spec["feature_start"]), int(spec["feature_end"]) + 1)
    features = [str(sheet.cell(1, column).value).strip() for column in columns]
    expected_features = 56 if kind == "mono" else 59
    if len(features) != expected_features or len(set(features)) != expected_features:
        raise RuntimeError(f"{kind}: missing or duplicate descriptor names")
    rows: list[dict[str, object]] = []
    for row_number in range(2, 26):
        sample = str(sheet.cell(row_number, int(spec["id_column"])).value).strip()
        expected_sample = f"{spec['prefix']}{row_number - 1}"
        if sample != expected_sample:
            raise RuntimeError(f"{kind}: expected {expected_sample}, found {sample}")
        lb = float(sheet.cell(row_number, int(spec["target_column"])).value)
        if not math.isfinite(lb) or lb <= 0:
            raise RuntimeError(f"{sample}: invalid L/B {lb}")
        source_name = str(sheet.cell(row_number, int(spec["source_name_column"])).value).strip()
        item: dict[str, object] = {"sample": sample, "source_name": source_name, "L_B": lb}
        for column, feature in zip(columns, features):
            value = float(sheet.cell(row_number, column).value)
            if not math.isfinite(value):
                raise RuntimeError(f"{sample}/{feature}: non-finite descriptor")
            item[feature] = value
        rows.append(item)
    return rows, features, sheet_name


def reference_feature_order(kind: str) -> list[str]:
    rows = read_dicts(Path(SOURCES[kind]["reference"]) / "feature_mapping.csv")
    return [row["source_descriptor"] for row in rows]


def inner_assignment(ligand_class: str, heldout: str, train_ids: list[str]) -> dict[str, int]:
    material = f"{FIXED_SEED}|{ligand_class}|{heldout}".encode("utf-8")
    seed = int.from_bytes(hashlib.sha256(material).digest()[:8], "big")
    ordered = sorted(train_ids)
    random.Random(seed).shuffle(ordered)
    return {ligand: index % 3 + 1 for index, ligand in enumerate(ordered)}


def write_snapshot(kind: str, rows: list[dict[str, object]], features: list[str]) -> None:
    fields = ["sample", "source_name", *features, "L_B"]
    write_dicts(ROOT / "input" / f"{kind}_full_data.csv", rows, fields)


def prepare_full_data(kind: str, rows: list[dict[str, object]], features: list[str]) -> None:
    spec = SOURCES[kind]
    target = str(spec["target"])
    selected_set, screening = screen(rows, features, target)
    selected = ordered_features(selected_set)
    expected = int(spec["expected_selected"])
    if len(selected) != expected:
        raise RuntimeError(f"{kind}: expected {expected} screened features, found {len(selected)}")
    reference_order = reference_feature_order(kind)
    if selected != reference_order:
        missing = sorted(set(reference_order) - set(selected))
        extra = sorted(set(selected) - set(reference_order))
        raise RuntimeError(f"{kind}: frozen feature mapping mismatch; missing={missing}, extra={extra}")
    case_dir = ROOT / "full_data" / kind
    write_dicts(case_dir / "screening.csv", screening, SCREENING_FIELDS)
    write_dicts(
        case_dir / "selected_features.csv",
        mapping_rows(selected, features),
        ["model_feature", "source_descriptor", "source_index", "ordering_family", "mono_unit_group"],
    )
    write_sisso_inputs(case_dir, rows, selected, features, kind, target, complexity=2)
    write_dicts(
        case_dir / "target_audit.csv",
        [
            {
                "sample": row["sample"],
                "raw_LB": format(float(row["L_B"]), ".17g"),
                "model_target": (
                    format(float(row["L_B"]), ".17g") if target == "raw"
                    else format(math.log(float(row["L_B"])), ".17g")
                ),
            }
            for row in rows
        ],
        ["sample", "raw_LB", "model_target"],
    )


def main() -> None:
    for name in ("cases", "full_data", "truth", "results"):
        reset_directory(ROOT / name)
    (ROOT / "slurm_logs").mkdir(exist_ok=True)

    loaded: dict[str, tuple[list[dict[str, object]], list[str]]] = {}
    source_manifest: list[dict[str, object]] = []
    for kind in ("mono", "bi"):
        rows, features, sheet_name = load_source(kind)
        loaded[kind] = (rows, features)
        write_snapshot(kind, rows, features)
        prepare_full_data(kind, rows, features)
        source_manifest.append({
            "ligand_class": kind,
            "workbook": Path(SOURCES[kind]["file"]).name,
            "sha256": SOURCES[kind]["sha256"],
            "worksheet": sheet_name,
            "samples": 24,
            "full_descriptor_count": len(features),
            "target": SOURCES[kind]["target"],
            "expected_full_data_screened_count": SOURCES[kind]["expected_selected"],
        })
    write_dicts(
        ROOT / "source_manifest.csv",
        source_manifest,
        ["ligand_class", "workbook", "sha256", "worksheet", "samples", "full_descriptor_count", "target", "expected_full_data_screened_count"],
    )

    pipelines = (
        ("M", "mono", "raw", loaded["mono"]),
        ("B", "bi", "log", loaded["bi"]),
    )
    manifest: list[dict[str, object]] = []
    truth: list[dict[str, object]] = []
    task_id = 0
    for pipeline, ligand_class, target, (rows, features) in pipelines:
        for outer_fold, heldout_row in enumerate(rows, start=1):
            task_id += 1
            heldout = str(heldout_row["sample"])
            case_rel = Path("cases") / f"{task_id:03d}_{pipeline}_{heldout}"
            case_dir = ROOT / case_rel
            case_dir.mkdir(parents=True)
            outer_training = [row for row in rows if row["sample"] != heldout]
            if len(outer_training) != 23:
                raise RuntimeError("Outer split did not contain 23 training ligands")
            fields = ["sample", "source_name", *features, "L_B"]
            write_dicts(case_dir / "outer_training.csv", outer_training, fields)
            write_dicts(
                case_dir / "heldout_features.csv",
                [{key: heldout_row[key] for key in ["sample", "source_name", *features]}],
                ["sample", "source_name", *features],
            )
            (case_dir / "outer_training_ligand_ids.txt").write_text(
                "\n".join(str(row["sample"]) for row in outer_training) + "\n", encoding="utf-8"
            )
            assignment = inner_assignment(
                ligand_class, heldout, [str(row["sample"]) for row in outer_training]
            )
            write_dicts(
                case_dir / "inner_splits.csv",
                [
                    {"ligand": row["sample"], "inner_validation_fold": assignment[str(row["sample"])]}
                    for row in outer_training
                ],
                ["ligand", "inner_validation_fold"],
            )
            metadata = {
                "task_id": task_id,
                "pipeline": pipeline,
                "ligand_class": ligand_class,
                "target": target,
                "outer_fold": outer_fold,
                "heldout_ligand": heldout,
                "n_train": 23,
                "fixed_seed": FIXED_SEED,
                "fixed_model": {"Phi": 2, "D": 4},
                "nested_grid": {"Phi": [1, 2], "D": [1, 2, 3, 4]},
                "operators": OPS,
                "safe_denominator": False,
            }
            (case_dir / "case_metadata.json").write_text(
                json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
            manifest.append({
                "task_id": task_id,
                "case_dir": case_rel.as_posix(),
                "pipeline": pipeline,
                "ligand_class": ligand_class,
                "target": target,
                "outer_fold": outer_fold,
                "heldout_ligand": heldout,
                "n_train": 23,
                "full_feature_count": len(features),
                "fixed_Phi": 2,
                "fixed_D": 4,
                "nested_Phi_grid": "1;2",
                "nested_D_grid": "1;2;3;4",
                "operators": OPS,
                "safe_denominator": False,
            })
            truth.append({
                "task_id": task_id,
                "pipeline": pipeline,
                "heldout_ligand": heldout,
                "L_B": format(float(heldout_row["L_B"]), ".17g"),
            })
    manifest_fields = [
        "task_id", "case_dir", "pipeline", "ligand_class", "target", "outer_fold",
        "heldout_ligand", "n_train", "full_feature_count", "fixed_Phi", "fixed_D",
        "nested_Phi_grid", "nested_D_grid", "operators", "safe_denominator",
    ]
    write_dicts(ROOT / "cases_manifest.csv", manifest, manifest_fields)
    write_dicts(ROOT / "truth" / "heldout_truth.csv", truth, ["task_id", "pipeline", "heldout_ligand", "L_B"])

    provenance = {
        "project": "LOOCV",
        "fixed_seed": FIXED_SEED,
        "operators": OPS,
        "safe_denominator": False,
        "source_workbook_sha256": {kind: SOURCES[kind]["sha256"] for kind in SOURCES},
        "reference_sha256": {
            f"{kind}/{path.name}": sha256(path)
            for kind in ("mono", "bi")
            for path in sorted(Path(SOURCES[kind]["reference"]).iterdir())
            if path.is_file()
        },
    }
    (ROOT / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=True, indent=2) + "\n", encoding="ascii"
    )
    print("PREPARE_OK: 2 full-data gates and 48 outer folds")
    print("OPERATORS=" + OPS)
    print("SAFE_DENOMINATOR=False")


if __name__ == "__main__":
    main()
